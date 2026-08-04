#!/usr/bin/env python3
"""Print the closing summary for ``make reproduce-experiment``.

Reports what was actually produced -- headline numbers straight out of the result JSONs,
the worksheet inventory of ``experimental_results.xlsx``, and the figure inventory -- and
returns a non-zero exit status if any promised artifact is absent or unpopulated. The
summary is a *check*, not just a printout: an AE checker should be able to trust that a
green ``make reproduce-experiment`` means every artifact really landed.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results"
FIGURES = ROOT / "figures"
XLSX = ROOT / "experimental_results.xlsx"

REQUIRED_SHEETS = [
    "Primary_Test_Results", "Adaptive_Attack_Results", "Per_Action_Breakdown",
    "Ablations_Analysis", "Baselines_Comparison", "Pilot_Decision_Gates",
]

# The 17 manuscript figures. The four step4_* diagnostics are supplementary and are
# reported separately rather than folded into this count.
MANUSCRIPT_FIGURES = [
    "fig01_graphical_abstract", "fig02_architecture", "fig03_security_utility_frontier",
    "fig04_per_action_profile", "fig05_ablation_effects",
    "fig06_verdicts_calibration_adaptive", "fig07_stage_attribution",
    "fig08_pilot_gates", "fig09_prevalence_regimes", "figA2_split_composition",
    "benchmark_corpus_composition", "benchmark_split_composition",
    "benchmark_authorization_structure", "benchmark_attack_design_balance",
    "step3_web_rag_environment", "step3_retrieval_profile", "step3_adaptive_attacker",
]
SUPPLEMENTARY_FIGURES = [
    "step4_security_utility_frontier", "step4_per_action_profile",
    "step4_ablation_effects", "step4_verdict_and_calibration",
]

REQUIRED_RESULTS = [
    "primary_test_evaluation.json", "adaptive_robustness_evaluation.json",
    "pilot_decision_report.json", "baseline_deviation_log.json",
    "evidence_family_identifiability.json", "defense_stage_traces_summary.json",
]


def rule(title: str = "") -> None:
    print("=" * 74 if not title else f"\n=== {title} " + "=" * max(0, 68 - len(title)))


def main() -> int:
    failures: list[str] = []
    rule()
    print(" reproduce-experiment SUMMARY")
    rule()

    # ------------------------------------------------------------------ results
    print("\nresults/")
    missing = []
    for name in REQUIRED_RESULTS:
        p = RESULTS / name
        if p.is_file():
            print(f"  OK      {name:<44} {p.stat().st_size / 1024:>9.0f} KB")
        else:
            missing.append(name)
            print(f"  MISSING {name}")
    if missing:
        failures.append(f"missing result artifacts: {missing}")
        return report(failures)

    prim = json.loads((RESULTS / "primary_test_evaluation.json").read_text())
    adapt = json.loads((RESULTS / "adaptive_robustness_evaluation.json").read_text())
    pilot = json.loads((RESULTS / "pilot_decision_report.json").read_text())

    # ---------------------------------------------------------------- headlines
    md = prim["metadata"]
    print("\nconfirmatory evaluation")
    print(f"  replay fingerprint verified : {md['replay_fingerprint_verified']}")
    print(f"  bootstrap resamples         : {md['n_resamples_used']} "
          f"(preregistered {md['preregistered_resamples']})")
    print(f"  confirmatory run            : {md['is_results_run']}")
    print(f"  configurations swept        : {md['n_configurations']}")
    print(f"  queries                     : {md['n_queries']}")
    print(f"  wall clock                  : {md['elapsed_seconds']:.1f} s")
    if not md["replay_fingerprint_verified"]:
        failures.append("replay fingerprint did NOT verify")
    if not md["is_results_run"]:
        failures.append(f"is_results_run is False (B = {md['n_resamples_used']} != "
                        f"{md['preregistered_resamples']}); this run is NOT confirmatory")
    if prim["frozen_parameters"].get("refit_on_evaluation_splits"):
        failures.append("refit_on_evaluation_splits is True; the run is not confirmatory")

    M, boot = prim["metrics"], prim["bootstrap"]["test"]
    ref = prim["reference_baseline_selection"]["reference_baseline"]
    aeg, rf = M["aegislink_full"]["test"], M[ref]["test"]

    def ci(cfg: str, metric: str) -> str:
        est = (boot.get("estimates") or {}).get(cfg, {}).get(metric, {})
        c = (est.get("ci") or {}).get("bias_corrected") or {}
        if c.get("lo") is None:
            return "[n/a]"
        tag = " (degenerate)" if c.get("degenerate") else ""
        return f"[{c['lo']:.4f}, {c['hi']:.4f}]{tag}"

    print("\nprimary endpoints on the confirmatory test split")
    print(f"  {'AegisLink (full)':<26} UALER {aeg['primary']['UALER']:.4f} "
          f"{ci('aegislink_full', 'UALER')}")
    print(f"  {'':<26} ATPR  {aeg['primary']['ATPR']:.4f} "
          f"{ci('aegislink_full', 'ATPR')}")
    print(f"  {'reference: ' + ref:<26} UALER {rf['primary']['UALER']:.4f} "
          f"{ci(ref, 'UALER')}")
    print(f"  {'':<26} ATPR  {rf['primary']['ATPR']:.4f} {ci(ref, 'ATPR')}")

    A = adapt["strata"]["A"]
    n_pages = sum(adapt["strata"][s]["replay"]["surface"]["n_pages"] for s in ("A", "B"))
    print(f"\nadaptive robustness ({n_pages} adaptive attack pages across 2 strata)")
    print(f"  stratum A (primary-claim eligible) ASR_a for AegisLink : "
          f"{A['asr']['aegislink_full']['ASR_a']:.4f} "
          f"({A['asr']['aegislink_full']['n_succeeded']}/"
          f"{A['asr']['aegislink_full']['n_attempted']})")
    print(f"  stratum A ASR_a for {ref:<24} : "
          f"{A['asr'][ref]['ASR_a']:.4f} "
          f"({A['asr'][ref]['n_succeeded']}/{A['asr'][ref]['n_attempted']})")

    d = pilot["decision"]
    print("\npilot Go/No-Go gates (CONTRACT.md Section 12)")
    print(f"  overall decision : {d['overall_decision']}  "
          f"(go_declared = {d['go_declared']})")
    print(f"  conditions       : {d['n_passed']} PASS, {d['n_failed']} FAIL, "
          f"{d['n_not_evaluable']} NOT_EVALUABLE of {d['n_conditions_declared']}")
    if d["not_evaluable_conditions"]:
        print(f"  not evaluable    : {', '.join(d['not_evaluable_conditions'])}")
    if d["failed_conditions"]:
        print(f"  failed           : {', '.join(d['failed_conditions'])}")

    # ------------------------------------------------------------------- excel
    print("\nexperimental_results.xlsx")
    if not XLSX.is_file():
        failures.append("experimental_results.xlsx was not produced")
        print("  MISSING")
    else:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(XLSX, read_only=True)
            for name in REQUIRED_SHEETS:
                if name not in wb.sheetnames:
                    failures.append(f"worksheet {name} is missing")
                    print(f"  MISSING {name}")
                    continue
                ws = wb[name]
                n_rows = ws.max_row or 0
                if n_rows < 2:
                    failures.append(f"worksheet {name} has no data rows")
                    print(f"  EMPTY   {name}")
                else:
                    print(f"  OK      {name:<28} {n_rows:>4} rows x "
                          f"{ws.max_column:>2} cols")
            extra = [s for s in wb.sheetnames if s not in REQUIRED_SHEETS]
            if extra:
                print(f"  plus supplementary sheet(s): {', '.join(extra)}")
            wb.close()
            print(f"  size    {XLSX.stat().st_size / 1024:.0f} KB")
        except ImportError:
            failures.append("openpyxl is unavailable, so the workbook cannot be verified")

    # ----------------------------------------------------------------- figures
    print("\nfigures/")
    have = {p.stem for p in FIGURES.glob("*.png")}
    miss_main = [f for f in MANUSCRIPT_FIGURES if f not in have]
    miss_supp = [f for f in SUPPLEMENTARY_FIGURES if f not in have]
    print(f"  manuscript figures    : {len(MANUSCRIPT_FIGURES) - len(miss_main)}"
          f"/{len(MANUSCRIPT_FIGURES)}")
    print(f"  supplementary figures : {len(SUPPLEMENTARY_FIGURES) - len(miss_supp)}"
          f"/{len(SUPPLEMENTARY_FIGURES)}")
    print(f"  total PNG in figures/ : {len(have)}   "
          f"(PDF companions: {len(list(FIGURES.glob('*.pdf')))})")
    if miss_main:
        failures.append(f"missing manuscript figures: {miss_main}")
        print(f"  MISSING {miss_main}")
    if miss_supp:
        failures.append(f"missing supplementary figures: {miss_supp}")
        print(f"  MISSING {miss_supp}")

    # ------------------------------------------------------------ honest scope
    print("\nwhat this reproduction does NOT cover (disclosed, not hidden)")
    for line in (
        "it replays a frozen retrieval snapshot; it does not re-crawl or re-index the web",
        "no language model is invoked anywhere -- the reader is a deterministic surrogate",
        "it does not run inside a pinned container image; docker/ is an empty contract slot",
        f"the pilot verdict is {d['overall_decision']}, not GO: "
        f"{d['n_not_evaluable']} condition(s) could not be measured by this pilot",
        "claim_ledger/ and manuscript/ are empty CONTRACT.md Section 17 slots in this bundle",
    ):
        print(f"  - {line}")

    return report(failures)


def report(failures: list[str]) -> int:
    print()
    rule()
    if failures:
        print(f" reproduce-experiment INCOMPLETE -- {len(failures)} problem(s)")
        for f in failures:
            print(f"   - {f}")
        rule()
        return 1
    print(" reproduce-experiment COMPLETE -- every promised artifact is present")
    rule()
    return 0


if __name__ == "__main__":
    sys.exit(main())
